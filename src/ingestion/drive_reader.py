import io
import logging
import os

from dotenv import load_dotenv
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

load_dotenv()

logger = logging.getLogger(__name__)

SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

GOOGLE_DOC_MIME = "application/vnd.google-apps.document"
GOOGLE_SHEET_MIME = "application/vnd.google-apps.spreadsheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def get_drive_service():
    creds_path = os.environ.get("GOOGLE_DRIVE_CREDENTIALS", "./service_account.json")
    creds = service_account.Credentials.from_service_account_file(creds_path, scopes=SCOPES)
    return build("drive", "v3", credentials=creds)


def _export_as(service, file_id: str, mime_type: str) -> bytes:
    request = service.files().export_media(fileId=file_id, mimeType=mime_type)
    buf = io.BytesIO()
    downloader = MediaIoBaseDownload(buf, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    return buf.getvalue()


def _download_binary(service, file_id: str) -> bytes:
    request = service.files().get_media(fileId=file_id)
    buf = io.BytesIO()
    downloader = MediaIoBaseDownload(buf, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    return buf.getvalue()


def _docx_to_text(file_bytes: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _xlsx_to_csv_text(file_bytes: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = ",".join(str(c) if c is not None else "" for c in row)
            if row_text.strip(","):
                lines.append(row_text)
    return "\n".join(lines)


def read_file_as_text(file_id: str, service=None) -> tuple[str, str, str]:
    """
    Read a Drive file and return (text, file_title, mime_type).
    Handles Google Docs, Google Sheets, .docx, .xlsx, and plain text files.
    """
    if service is None:
        service = get_drive_service()

    try:
        meta = service.files().get(fileId=file_id, fields="id,name,mimeType,size").execute()
        file_title = meta["name"]
        mime_type = meta["mimeType"]

        if mime_type == GOOGLE_DOC_MIME:
            raw = _export_as(service, file_id, "text/plain")
            text = raw.decode("utf-8", errors="replace")

        elif mime_type == GOOGLE_SHEET_MIME:
            raw = _export_as(service, file_id, "text/csv")
            text = raw.decode("utf-8", errors="replace")

        elif mime_type == DOCX_MIME:
            file_bytes = _download_binary(service, file_id)
            text = _docx_to_text(file_bytes)

        elif mime_type == XLSX_MIME:
            # Try Drive export to CSV first; fall back to openpyxl
            try:
                raw = _export_as(service, file_id, "text/csv")
                text = raw.decode("utf-8", errors="replace")
            except Exception:
                file_bytes = _download_binary(service, file_id)
                text = _xlsx_to_csv_text(file_bytes)

        else:
            # Best-effort: try text export, then raw download
            try:
                raw = _export_as(service, file_id, "text/plain")
                text = raw.decode("utf-8", errors="replace")
            except Exception:
                file_bytes = _download_binary(service, file_id)
                text = file_bytes.decode("utf-8", errors="replace")

        logger.info("Read '%s' (%s): %d chars", file_title, mime_type, len(text))
        return text, file_title, mime_type

    except Exception as e:
        raise RuntimeError(f"Failed to read Drive file {file_id}: {e}") from e
